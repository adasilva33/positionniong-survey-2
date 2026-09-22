"""
Positioning Survey — automated Google Form generation.

Reads `positioning_survey_universe.xlsx` and creates ONE Google Form per Sector.
Page 1 has the consent text plus BROKER / NAME / EMAIL. Then one section (page
break) per Subsector, headed "SECTOR: SUBSECTOR", each with 9 dropdown questions
(CONSENSUS LONG / SHORT / CONTENTIOUS x TICKER 1-3), every dropdown populated with
that subsector's tickers only. Each run creates fresh forms (v1 — no update yet).

Full setup + per-cycle checklist: see GUIDE.md. Capacity limits: CAPACITY_MATRIX.md.
--------------------------------------------------------------------------------
Setup
  1. pip install google-api-python-client google-auth-httplib2 \
                 google-auth-oauthlib openpyxl
  2. Put the OAuth *Desktop app* client `credentials.json` next to this script.
  3. First run opens a browser for consent; a `token.json` cache is written after.

Titles, page text, identification fields and question groups are all set in
`survey_config.toml` (edit that, not this file). It is optional — built-in
defaults apply if it is absent.

Usage
  # Real run: workbook + survey_config.toml from the current directory
  python create_form.py

  # Point at a different workbook / config
  python create_form.py --input path/to/universe.xlsx --config path/to/config.toml

  # See exactly what WOULD be created — no browser, no API calls
  python create_form.py --dry-run

  # Only build the first N sectors (smoke-test one real form from a big sheet)
  python create_form.py --limit 1

  # Generate a synthetic workbook to stress-test scale, then dry-run it
  python create_form.py --make-sample sample.xlsx --sectors 100 --subsectors 10 --tickers 40
  python create_form.py --input sample.xlsx --dry-run
--------------------------------------------------------------------------------
"""

import argparse
import csv
import os
import sys
import time
import tomllib
from collections import OrderedDict

SCOPES = ["https://www.googleapis.com/auth/forms.body"]
CREDENTIALS_FILE = "credentials.json"
TOKEN_FILE = "token.json"

UNIVERSE_SHEET = "Universe"
INTRO_SHEET = "Intro Text"
DEFAULT_INPUT = "positioning_survey_universe.xlsx"
DEFAULT_CONFIG = "survey_config.toml"
OUTPUT_CSV = "created_forms.csv"

# ---- Defaults. survey_config.toml overrides these; then CLI flags override that.
# Form title template. Placeholders: {region} {sector} {period}. Upper-cased.
TITLE_TEMPLATE = "{region} {sector} POSITIONING SURVEY - {period}"
TITLE_OVERRIDES = {}          # {sector: exact title}; used verbatim
SURVEY_REGION = "EMEA"          # prefix; set "" to drop it
SURVEY_PERIOD = ""             # e.g. "Q2 2026"; blank -> current calendar quarter

# NOTE: the form's header/logo banner is NOT settable via the Forms API (UI-only,
# per form). Set it by hand after creation, or clone a pre-branded template form.

# Fallback description if a sector has no row in the `Intro Text` sheet.
DEFAULT_INTRO_TEXT = (
    "We may consider the information that you provide when deciding whether to buy "
    "or sell stock or other investments. Therefore, we do not want to receive "
    "material, non-public information or information that is subject to a "
    "confidentiality obligation."
)

# Appended under the consent text on page 1 (mirrors the MS Forms notice).
SUBMIT_NOTICE = (
    "When you submit this form, it will not automatically collect your details "
    "like name and email address unless you provide it yourself."
)

# First page of every form: respondent identification. (title, help text).
# All required short-answer questions, shown before the first subsector section.
IDENTIFICATION_QUESTIONS = [
    ("BROKER", ""),
    ("NAME OF PERSON COMPLETING", ""),
    ("EMAIL ADDRESS", "Please enter your work email address."),
]

# The 9 questions per subsector: (group label, ticker count). Ranking is
# explained once in the section description, not repeated in every title.
QUESTION_GROUPS = [
    ("CONSENSUS LONG", 3),
    ("CONSENSUS SHORT", 3),
    ("CONTENTIOUS", 3),
]
QUESTIONS_PER_SUBSECTOR = sum(n for _, n in QUESTION_GROUPS)  # 9

# Shown as the description under every "SECTOR: SUBSECTOR" section header.
SECTION_INSTRUCTIONS = (
    "Please enter your top 3 BBG tickers for each category (ranked).\n"
    "- Consensus Longs: Ticker 1 = strongest long, Ticker 2 = second, Ticker 3 = third\n"
    "- Consensus Shorts: Ticker 1 = strongest short, Ticker 2 = second, Ticker 3 = third\n"
    "- Contentious: Ticker 1-3 = debated names"
)

# Split each form's item creation into batchUpdate calls of this many requests.
# Larger = more likely the whole form lands in one atomic call (a mid-way failure
# otherwise leaves a partially-built form, which the API can't delete).
CHUNK_SIZE = 200
# Options past this in a flat dropdown get unwieldy (no type-ahead in Google Forms).
DROPDOWN_SOFT_LIMIT = 100
# Google Forms rejects a form once its TOTAL option count across all questions
# gets too large ("exceeding the choice limit"). The exact number is undocumented;
# this is a conservative pre-flight ceiling. Override with `max_total_choices` in
# the config or --force to try anyway. Truncate per dropdown with --max-options.
MAX_TOTAL_CHOICES = 4000
MAX_OPTIONS_PER_DROPDOWN = 0  # 0 = no cap
# Seconds to pause between forms, to stay well under the Forms API write quota.
SLEEP_BETWEEN_FORMS = 1.0


def survey_period():
    if SURVEY_PERIOD:
        return SURVEY_PERIOD
    now = time.localtime()
    return f"Q{(now.tm_mon - 1) // 3 + 1} {now.tm_year}"


def form_title(sector):
    if sector in TITLE_OVERRIDES:
        return TITLE_OVERRIDES[sector]
    t = TITLE_TEMPLATE.format(
        region=SURVEY_REGION, sector=sector, period=survey_period()
    )
    return " ".join(t.split()).upper()  # collapse the gap left by an empty {region}


def load_config(path):
    """Overlay survey_config.toml onto the module-level defaults. Missing file
    is fine — built-in defaults stand."""
    global TITLE_TEMPLATE, TITLE_OVERRIDES, SURVEY_REGION, SURVEY_PERIOD
    global DEFAULT_INTRO_TEXT, SUBMIT_NOTICE, SECTION_INSTRUCTIONS
    global IDENTIFICATION_QUESTIONS, QUESTION_GROUPS, QUESTIONS_PER_SUBSECTOR
    global MAX_TOTAL_CHOICES, MAX_OPTIONS_PER_DROPDOWN

    if not os.path.exists(path):
        if path != DEFAULT_CONFIG:
            sys.exit(f"Config file not found: {path}")
        return
    with open(path, "rb") as f:
        cfg = tomllib.load(f)

    TITLE_TEMPLATE = cfg.get("title", TITLE_TEMPLATE)
    TITLE_OVERRIDES = dict(cfg.get("title_overrides", {}))
    SURVEY_REGION = cfg.get("region", SURVEY_REGION)
    SURVEY_PERIOD = cfg.get("period", SURVEY_PERIOD)
    DEFAULT_INTRO_TEXT = cfg.get("intro_text", DEFAULT_INTRO_TEXT)
    SUBMIT_NOTICE = cfg.get("submit_notice", SUBMIT_NOTICE)
    SECTION_INSTRUCTIONS = cfg.get("section_instructions", SECTION_INSTRUCTIONS)

    ident = cfg.get("identification")
    if ident:
        IDENTIFICATION_QUESTIONS = [(q["title"], q.get("help", "")) for q in ident]
    groups = cfg.get("questions")
    if groups:
        QUESTION_GROUPS = [(label, int(n)) for label, n in groups.items()]
        QUESTIONS_PER_SUBSECTOR = sum(n for _, n in QUESTION_GROUPS)

    limits = cfg.get("limits", {})
    MAX_TOTAL_CHOICES = int(limits.get("max_total_choices", MAX_TOTAL_CHOICES))
    MAX_OPTIONS_PER_DROPDOWN = int(
        limits.get("max_options_per_dropdown", MAX_OPTIONS_PER_DROPDOWN)
    )


# --------------------------------------------------------------------------- #
# Auth
# --------------------------------------------------------------------------- #
def get_forms_service():
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    creds = None
    if os.path.exists(TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not os.path.exists(CREDENTIALS_FILE):
                sys.exit(
                    f"Missing {CREDENTIALS_FILE}. Download the OAuth Desktop-app "
                    "client from Google Cloud Console and put it next to this script."
                )
            flow = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(TOKEN_FILE, "w") as f:
            f.write(creds.to_json())
    return build("forms", "v1", credentials=creds)


def execute_with_retry(request, tries=5):
    """Run an API request, backing off on 429/5xx."""
    from googleapiclient.errors import HttpError

    delay = 2.0
    for attempt in range(1, tries + 1):
        try:
            return request.execute()
        except HttpError as e:
            status = getattr(e.resp, "status", None)
            if status in (429, 500, 502, 503) and attempt < tries:
                print(f"  API {status}; retry {attempt}/{tries - 1} in {delay:.0f}s")
                time.sleep(delay)
                delay *= 2
            else:
                raise


# --------------------------------------------------------------------------- #
# Workbook loading
# --------------------------------------------------------------------------- #
def _norm(v):
    return str(v).strip() if v is not None else ""


def load_universe(path):
    """Return (sectors, intros).

    sectors : OrderedDict[sector] -> OrderedDict[subsector] -> [ticker, ...]
    intros  : dict[sector] -> description text
    Order of first appearance in the sheet is preserved throughout.
    """
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl not installed. Run: pip install openpyxl")

    if not os.path.exists(path):
        sys.exit(f"Input workbook not found: {path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    if UNIVERSE_SHEET not in wb.sheetnames:
        sys.exit(f"Workbook has no '{UNIVERSE_SHEET}' sheet. Found: {wb.sheetnames}")

    ws = wb[UNIVERSE_SHEET]
    rows = ws.iter_rows(values_only=True)
    header = [_norm(c).lower() for c in next(rows, [])]
    try:
        i_sec = header.index("sector")
        i_sub = header.index("subsector")
        i_tkr = header.index("ticker")
    except ValueError:
        sys.exit(
            f"'{UNIVERSE_SHEET}' must have columns: Sector | Subsector | Ticker "
            f"(found: {header})"
        )

    sectors = OrderedDict()
    seen_pairs = set()  # (sector, subsector, ticker) dedupe
    for raw in rows:
        if raw is None:
            continue
        sector = _norm(raw[i_sec]) if i_sec < len(raw) else ""
        subsector = _norm(raw[i_sub]) if i_sub < len(raw) else ""
        ticker = _norm(raw[i_tkr]) if i_tkr < len(raw) else ""
        # Skip blank rows and the trailing legend/notes block.
        if not (sector and subsector and ticker):
            continue
        key = (sector, subsector, ticker)
        if key in seen_pairs:
            continue
        seen_pairs.add(key)
        sectors.setdefault(sector, OrderedDict()).setdefault(subsector, []).append(ticker)

    if not sectors:
        sys.exit(f"No usable rows in '{UNIVERSE_SHEET}'.")

    intros = {}
    if INTRO_SHEET in wb.sheetnames:
        iws = wb[INTRO_SHEET]
        irows = iws.iter_rows(values_only=True)
        ihdr = [_norm(c).lower() for c in next(irows, [])]
        si = ihdr.index("sector") if "sector" in ihdr else 0
        # The text column: anything that isn't the sector column.
        ti = next((j for j, h in enumerate(ihdr) if h and j != si), 1)
        for raw in irows:
            if not raw:
                continue
            sec = _norm(raw[si]) if si < len(raw) else ""
            txt = _norm(raw[ti]) if ti < len(raw) else ""
            if sec and txt:
                intros[sec] = txt

    wb.close()
    return sectors, intros


# --------------------------------------------------------------------------- #
# Request building
# --------------------------------------------------------------------------- #
def build_form_requests(sector, subsectors, intro_text):
    """All batchUpdate requests for one sector's form.

    `subsectors` : OrderedDict[subsector] -> [ticker, ...]
    Item indices are assigned sequentially from 0; sending the resulting list in
    order (even split across several batchUpdate calls) inserts correctly because
    every createItem appends at the current tail.
    """
    requests = []
    description = "\n\n".join(p for p in (intro_text, SUBMIT_NOTICE) if p)
    if description:
        requests.append(
            {
                "updateFormInfo": {
                    "info": {"description": description},
                    "updateMask": "description",
                }
            }
        )

    idx = 0
    # Respondent identification questions on the first page (before any page break).
    for q_title, q_help in IDENTIFICATION_QUESTIONS:
        item = {
            "title": q_title,
            "questionItem": {
                "question": {
                    "required": True,
                    "textQuestion": {"paragraph": False},
                }
            },
        }
        if q_help:
            item["description"] = q_help
        requests.append({"createItem": {"item": item, "location": {"index": idx}}})
        idx += 1

    for subsector, tickers in subsectors.items():
        if MAX_OPTIONS_PER_DROPDOWN:
            tickers = tickers[:MAX_OPTIONS_PER_DROPDOWN]
        options = [{"value": t} for t in tickers]
        requests.append(
            {
                "createItem": {
                    "item": {
                        "title": f"{sector}: {subsector}".upper(),
                        "description": SECTION_INSTRUCTIONS,
                        "pageBreakItem": {},
                    },
                    "location": {"index": idx},
                }
            }
        )
        idx += 1
        for group_label, n_ranks in QUESTION_GROUPS:
            for rank in range(1, n_ranks + 1):
                requests.append(
                    {
                        "createItem": {
                            "item": {
                                "title": f"{group_label}: TICKER {rank}",
                                "questionItem": {
                                    "question": {
                                        "required": False,
                                        "choiceQuestion": {
                                            "type": "DROP_DOWN",
                                            "options": options,
                                        },
                                    }
                                },
                            },
                            "location": {"index": idx},
                        }
                    }
                )
                idx += 1
    return requests


def chunked(seq, n):
    for i in range(0, len(seq), n):
        yield seq[i : i + n]


# --------------------------------------------------------------------------- #
# Planning / validation
# --------------------------------------------------------------------------- #
def _capped(tickers):
    return tickers[:MAX_OPTIONS_PER_DROPDOWN] if MAX_OPTIONS_PER_DROPDOWN else tickers


def plan_rows(sectors):
    """Yield dicts describing each sector for reporting/validation."""
    for sector, subs in sectors.items():
        n_sub = len(subs)
        n_tkr = sum(len(v) for v in subs.values())
        items = (
            len(IDENTIFICATION_QUESTIONS)  # broker / name / email
            + n_sub * (1 + QUESTIONS_PER_SUBSECTOR)
        )
        requests = items + 1  # + updateFormInfo (description)
        choices = sum(len(_capped(v)) for v in subs.values()) * QUESTIONS_PER_SUBSECTOR
        yield {
            "sector": sector,
            "subsectors": n_sub,
            "tickers": n_tkr,
            "questions": n_sub * QUESTIONS_PER_SUBSECTOR,
            "items": items,
            "choices": choices,
            "over_limit": choices > MAX_TOTAL_CHOICES,
            "batch_calls": max(1, -(-requests // CHUNK_SIZE)),  # ceil
        }


def validate(sectors):
    warnings = []
    for sector, subs in sectors.items():
        for subsector, tickers in subs.items():
            if len(tickers) < 3:
                warnings.append(
                    f"{sector} / {subsector}: only {len(tickers)} ticker(s) — "
                    "dropdowns for rank 2/3 will be sparse."
                )
            if len(tickers) > DROPDOWN_SOFT_LIMIT:
                warnings.append(
                    f"{sector} / {subsector}: {len(tickers)} tickers — past the "
                    f"~{DROPDOWN_SOFT_LIMIT} usability limit for a flat dropdown."
                )
    for r in plan_rows(sectors):
        if r["over_limit"]:
            warnings.append(
                f"{r['sector']}: ~{r['choices']} total dropdown options "
                f"(> {MAX_TOTAL_CHOICES}) — Google Forms will reject this form. "
                f"Reduce tickers/subsector, set max_options_per_dropdown, or --force."
            )
    return warnings


def print_plan(sectors, intros):
    rows = list(plan_rows(sectors))
    tot_sub = sum(r["subsectors"] for r in rows)
    tot_q = sum(r["questions"] for r in rows)
    tot_calls = sum(r["batch_calls"] + 1 for r in rows)  # +1 create() per form

    print(f"\nUniverse: {len(rows)} sector(s), {tot_sub} subsector(s), "
          f"{sum(r['tickers'] for r in rows)} ticker rows")
    print(f"Title    : {form_title(rows[0]['sector']) if rows else '-'}")
    print(f"{'SECTOR':<24}{'SUBSECS':>8}{'TICKERS':>9}{'QS':>6}{'CHOICES':>9}"
          f"{'BATCHES':>9}  INTRO")
    print("-" * 80)
    for r in rows:
        intro = "sheet" if r["sector"] in intros else "default"
        flag = "  !! over choice limit" if r["over_limit"] else ""
        print(f"{r['sector'][:23]:<24}{r['subsectors']:>8}{r['tickers']:>9}"
              f"{r['questions']:>6}{r['choices']:>9}{r['batch_calls']:>9}  {intro}{flag}")
    print("-" * 80)
    print(f"{'TOTAL':<24}{tot_sub:>8}{'':>9}{tot_q:>6}"
          f"{sum(r['choices'] for r in rows):>9}")
    print(f"\nApprox API calls if run for real: ~{tot_calls} "
          f"({len(rows)} forms.create + {tot_calls - len(rows)} batchUpdate)")

    warns = validate(sectors)
    if warns:
        print(f"\n{len(warns)} warning(s):")
        for w in warns:
            print(f"  ! {w}")
    else:
        print("\nNo validation warnings.")


# --------------------------------------------------------------------------- #
# Sample workbook generator (for scale-testing)
# --------------------------------------------------------------------------- #
def make_sample(path, n_sectors, n_subsectors, n_tickers):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = UNIVERSE_SHEET
    ws.append(["Sector", "Subsector", "Ticker"])
    for s in range(1, n_sectors + 1):
        sector = f"Sector {s}"
        for u in range(1, n_subsectors + 1):
            subsector = f"Subsector {u}"
            for t in range(1, n_tickers + 1):
                ws.append([sector, subsector, f"{sector} {subsector} Ticker {t}"])

    iws = wb.create_sheet(INTRO_SHEET)
    iws.append(["Sector", "Intro / Description Text"])
    for s in range(1, n_sectors + 1):
        iws.append([f"Sector {s}", DEFAULT_INTRO_TEXT])

    wb.save(path)
    total = n_sectors * n_subsectors * n_tickers
    print(f"Wrote {path}: {n_sectors} sectors x {n_subsectors} subsectors x "
          f"{n_tickers} tickers = {total} rows")


# --------------------------------------------------------------------------- #
# Real creation
# --------------------------------------------------------------------------- #
def create_forms(service, sectors, intros, limit=None, force=False):
    from googleapiclient.errors import HttpError

    results = []
    skipped = []
    items = list(sectors.items())
    if limit:
        items = items[:limit]
    plan = {r["sector"]: r for r in plan_rows(dict(items))}

    for n, (sector, subs) in enumerate(items, 1):
        title = form_title(sector)
        r = plan[sector]
        print(f"\n[{n}/{len(items)}] {sector}: {len(subs)} subsectors, "
              f"{r['tickers']} tickers, ~{r['choices']} dropdown options")

        if r["over_limit"] and not force:
            print(f"  SKIPPED: ~{r['choices']} options exceeds MAX_TOTAL_CHOICES "
                  f"({MAX_TOTAL_CHOICES}). Trim the universe, set "
                  f"max_options_per_dropdown, or re-run with --force.")
            skipped.append(sector)
            continue

        form = execute_with_retry(
            service.forms().create(
                body={"info": {"title": title, "documentTitle": title}}
            )
        )
        form_id = form["formId"]
        edit_url = f"https://docs.google.com/forms/d/{form_id}/edit"
        responder_url = form.get("responderUri", "")

        requests = build_form_requests(
            sector, subs, intros.get(sector, DEFAULT_INTRO_TEXT)
        )
        n_chunks = max(1, -(-len(requests) // CHUNK_SIZE))
        try:
            for c, chunk in enumerate(chunked(requests, CHUNK_SIZE), 1):
                print(f"  batchUpdate {c}/{n_chunks} ({len(chunk)} requests)")
                execute_with_retry(
                    service.forms().batchUpdate(
                        formId=form_id, body={"requests": chunk}
                    )
                )
        except HttpError as e:
            msg = str(e)
            if "choice limit" in msg or "entry limit" in msg:
                print(f"  FAILED: Google Forms rejected the form — {msg.splitlines()[0]}")
                print(f"  Earlier batches may have committed — this form is likely "
                      f"partially built. The Forms API cannot delete forms; trash it "
                      f"in Drive: {edit_url}")
            else:
                print(f"  FAILED: {msg.splitlines()[0]}")
            skipped.append(sector)
            continue

        print(f"  done -> {edit_url}")
        results.append(
            {
                "sector": sector,
                "form_id": form_id,
                "edit_url": edit_url,
                "responder_url": responder_url,
                "subsectors": len(subs),
                "questions": len(subs) * QUESTIONS_PER_SUBSECTOR,
            }
        )
        if n < len(items):
            time.sleep(SLEEP_BETWEEN_FORMS)

    if skipped:
        print(f"\n{len(skipped)} sector(s) not created: {', '.join(skipped)}")

    if results:
        with open(OUTPUT_CSV, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(results[0].keys()))
            w.writeheader()
            w.writerows(results)
        print(f"\nWrote {OUTPUT_CSV} ({len(results)} form(s)).")
    return results


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main():
    global SURVEY_REGION, SURVEY_PERIOD, MAX_OPTIONS_PER_DROPDOWN
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--input", default=DEFAULT_INPUT, help=f"universe workbook (default: {DEFAULT_INPUT})")
    p.add_argument("--dry-run", action="store_true", help="print the plan only; no auth, no API calls")
    p.add_argument("--limit", type=int, metavar="N", help="only build the first N sectors")
    p.add_argument("--make-sample", metavar="PATH", help="write a synthetic workbook and exit")
    p.add_argument("--sectors", type=int, default=2, help="--make-sample: sector count (default 2)")
    p.add_argument("--subsectors", type=int, default=10, help="--make-sample: subsectors per sector (default 10)")
    p.add_argument("--tickers", type=int, default=10, help="--make-sample: tickers per subsector (default 10)")
    p.add_argument("--config", default=DEFAULT_CONFIG, help=f"TOML config (default: {DEFAULT_CONFIG})")
    p.add_argument("--region", metavar="STR", help='title region prefix ("" to drop)')
    p.add_argument("--period", metavar="STR", help='title period, e.g. "Q2 2026" (default: current quarter)')
    p.add_argument("--max-options", type=int, metavar="N", help="truncate every dropdown to N options (0 = no cap)")
    p.add_argument("--force", action="store_true", help="create forms even if they exceed the choice-limit estimate")
    args = p.parse_args()

    load_config(args.config)

    if args.region is not None:
        SURVEY_REGION = args.region
    if args.period is not None:
        SURVEY_PERIOD = args.period
    if args.max_options is not None:
        MAX_OPTIONS_PER_DROPDOWN = args.max_options

    if args.make_sample:
        make_sample(args.make_sample, args.sectors, args.subsectors, args.tickers)
        return

    sectors, intros = load_universe(args.input)

    if args.dry_run:
        print_plan(sectors, intros)
        print("\nDry run — nothing created.")
        return

    print_plan(sectors, intros)
    service = get_forms_service()
    create_forms(service, sectors, intros, limit=args.limit, force=args.force)

    print(
        "\nReminder: per form, in the Forms UI set "
        "Settings -> Responses -> 'Limit to 1 response' OFF and email collection "
        "OFF so external respondents are not forced to sign in. This is not "
        "exposed by the API."
    )


if __name__ == "__main__":
    main()
